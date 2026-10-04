"""Shell-friendly copies of files the hosted shell cannot open by itself."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import PurePosixPath
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

#: Office Open XML workbooks; legacy ``.xls`` needs a reader nobody ships.
_WORKBOOK_SUFFIXES = frozenset({".xlsx", ".xlsm"})


@dataclass(frozen=True, slots=True)
class Rendition:
    """One derived file uploaded beside the original."""

    file_name: str
    mime_type: str
    data: bytes


def renditions(file_name: str, data: bytes) -> tuple[Rendition, ...]:
    """Return the copies the shell needs to read ``file_name``, if any.

    The hosted shell has pandas but no Excel reader and no internet, so a
    workbook also arrives as one CSV per non-empty sheet. Every cell value is
    kept as is, header rows included: deciding which row is the header is the
    analysis, not the copy. A workbook that will not open yields no copies and
    the original is still delivered.
    """

    path = PurePosixPath(file_name)
    if path.suffix.lower() not in _WORKBOOK_SUFFIXES:
        return ()
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        return ()
    try:
        sheets = [
            (sheet.title, _csv(sheet.iter_rows(values_only=True)))
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()
    sheets = [(title, text) for title, text in sheets if text]
    single = len(sheets) == 1
    return tuple(
        Rendition(
            file_name=f"{path.stem}.csv" if single else f"{path.stem}.{_safe(title)}.csv",
            mime_type="text/csv",
            data=text.encode("utf-8"),
        )
        for title, text in sheets
    )


def _csv(rows: object) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    pending_blank = 0
    wrote = False
    for row in rows:  # type: ignore[attr-defined]
        values = [_cell(value) for value in row]
        while values and values[-1] == "":
            values.pop()
        if not values:
            pending_blank += 1
            continue
        # Blank rows inside the data stay; trailing ones are dropped.
        for _ in range(pending_blank if wrote else 0):
            writer.writerow(())
        pending_blank = 0
        writer.writerow(values)
        wrote = True
    return buffer.getvalue()


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _safe(title: str) -> str:
    name = re.sub(r"[^\w.-]+", "_", title, flags=re.UNICODE).strip("._")
    return name or "sheet"


__all__ = ["Rendition", "renditions"]
